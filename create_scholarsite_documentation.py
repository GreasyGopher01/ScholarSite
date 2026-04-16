from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

doc.add_heading('ScholarSite Technical Documentation', level=0)

p = doc.add_paragraph(
    'Generated technical documentation for the ScholarSite application. This document covers architecture, setup, backend and frontend details, data model, authentication flow, deployment, and testing.'
)
p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

doc.add_heading('1. Overview', level=1)
doc.add_paragraph(
    'ScholarSite is a full-stack application for discovering, bookmarking, and recommending scholarship and opportunity resources. It combines a FastAPI backend with a MongoDB datastore and an Angular frontend. The application uses Google authentication and personalized recommendation logic based on bookmarked opportunities.'
)

doc.add_heading('2. Architecture', level=1)
doc.add_paragraph(
    'The application follows a decoupled architecture with separate frontend and backend layers.'
)
doc.add_paragraph('Backend: FastAPI API server written in Python.', style='List Bullet')
doc.add_paragraph('Frontend: Angular 21 standalone components with client-side routing.', style='List Bullet')
doc.add_paragraph('Database: MongoDB for persistence of users, bookmarks, opportunities, tips, feedback, and subscriptions.', style='List Bullet')

doc.add_heading('2.1 Backend Responsibilities', level=2)
doc.add_paragraph('The backend handles:')
doc.add_paragraph('User authentication and JWT token issuance', style='List Bullet')
doc.add_paragraph('CRUD operations for bookmarks, feedback, subscriptions, and profile updates', style='List Bullet')
doc.add_paragraph('Serving opportunity and tip data from MongoDB', style='List Bullet')
doc.add_paragraph('CORS configuration for frontend origins', style='List Bullet')

doc.add_heading('2.2 Frontend Responsibilities', level=2)
doc.add_paragraph('The frontend handles:')
doc.add_paragraph('Google sign-in integration and redirect flow', style='List Bullet')
doc.add_paragraph('Rendering opportunity directory, bookmarks, recommendations, and profile pages', style='List Bullet')
doc.add_paragraph('Communicating with backend APIs via HTTP client', style='List Bullet')
doc.add_paragraph('Local user session persistence via browser localStorage', style='List Bullet')

doc.add_heading('3. Backend', level=1)
doc.add_paragraph(
    'The backend is implemented in Python using FastAPI and Motor to communicate with MongoDB asynchronously. Core backend features include authentication, user profile management, bookmarks, opportunity retrieval, tips, recommendations, feedback, and subscription management.'
)

doc.add_heading('3.1 Key Backend Files', level=2)
doc.add_paragraph('• api/main.py: Main FastAPI application and endpoint definitions.')
doc.add_paragraph('• api/requirements.txt: Python dependency manifest.')
doc.add_paragraph('• api/scripts/: Data seeding scripts for opportunities and tips.')

doc.add_heading('3.2 Authentication Flow', level=2)
doc.add_paragraph('Authentication uses Google OAuth tokens and JWTs:')
doc.add_paragraph('1. The frontend obtains a Google credential using the Google identity SDK.', style='List Number')
doc.add_paragraph('2. The credential is posted to /api/auth/google in the backend.', style='List Number')
doc.add_paragraph('3. The backend verifies the Google token with google-auth.', style='List Number')
doc.add_paragraph('4. If the user is new, a MongoDB user record is created. Otherwise the record is updated.', style='List Number')
doc.add_paragraph('5. A JWT is generated and returned to the frontend with the user object.', style='List Number')

doc.add_heading('3.3 API Endpoints', level=2)
doc.add_paragraph('Selected backend endpoints:')
doc.add_paragraph('• GET /: Health and API metadata', style='List Bullet')
doc.add_paragraph('• POST /api/auth/google: Google login', style='List Bullet')
doc.add_paragraph('• GET /api/auth/me: Verify current JWT and return user info', style='List Bullet')
doc.add_paragraph('• PUT /api/auth/profile: Update user profile fields', style='List Bullet')
doc.add_paragraph('• GET /opportunities: Retrieve active opportunities', style='List Bullet')
doc.add_paragraph('• GET /tips: Retrieve tip documents', style='List Bullet')
doc.add_paragraph('• GET /api/bookmarks: Get the current user bookmarks', style='List Bullet')
doc.add_paragraph('• POST /api/bookmarks: Add a bookmark for the current user', style='List Bullet')
doc.add_paragraph('• DELETE /api/bookmarks/{id}: Remove a bookmark', style='List Bullet')

doc.add_heading('3.4 Data Stores', level=2)
doc.add_paragraph('Collections used in MongoDB:')
doc.add_paragraph('• users: Stores authenticated user data, profile fields, and provider metadata.', style='List Bullet')
doc.add_paragraph('• bookmarks: Stores userId and opportunityId associations.', style='List Bullet')
doc.add_paragraph('• opportunities: Stores opportunity records with source links, categories, location, and deadlines.', style='List Bullet')
doc.add_paragraph('• tips: Stores tip content items.', style='List Bullet')
doc.add_paragraph('• feedback: Stores user feedback messages.', style='List Bullet')
doc.add_paragraph('• subscriptions: Stores notification subscription preferences.', style='List Bullet')

doc.add_heading('4. Frontend', level=1)
doc.add_paragraph(
    'The frontend is an Angular application with standalone components, client-side routing, HTTP access, and reactive forms. It is located in the ui/ directory.'
)

doc.add_heading('4.1 Key Frontend Files', level=2)
doc.add_paragraph('• ui/src/app/services/auth.service.ts: Authentication service, Google login, profile updates, bookmarks API wrappers.', style='List Bullet')
doc.add_paragraph('• ui/src/app/components/login/login.component.ts: Login flow initialization and Google sign-in rendering.', style='List Bullet')
doc.add_paragraph('• ui/src/app/components/directory/directory.component.ts: Opportunity directory browsing and filtering.', style='List Bullet')
doc.add_paragraph('• ui/src/app/components/recommendations/recommendations.component.ts: Personalized recommendations logic and scoring.', style='List Bullet')
doc.add_paragraph('• ui/src/app/components/bookmarks/bookmarks.component.ts: Bookmark management UI.', style='List Bullet')
doc.add_paragraph('• ui/src/environments/environment.prod.ts: Production API endpoint configuration.', style='List Bullet')

doc.add_heading('4.2 Routing', level=2)
doc.add_paragraph('Routes are defined in ui/src/app/app.routes.ts and include:')
doc.add_paragraph('• /login', style='List Bullet')
doc.add_paragraph('• /directory', style='List Bullet')
doc.add_paragraph('• /bookmarks', style='List Bullet')
doc.add_paragraph('• /recommendations', style='List Bullet')
doc.add_paragraph('• /tips', style='List Bullet')
doc.add_paragraph('• /feedback', style='List Bullet')

doc.add_heading('4.3 Recommendation Algorithm', level=2)
doc.add_paragraph('Recommendations are generated based on the current user bookmarks. The algorithm analyzes bookmarked opportunity categories, states, keywords, and tags, then scores candidate opportunities by similarity and returns the top results.')

doc.add_heading('4.4 UI Behavior', level=2)
doc.add_paragraph('• Bookmarks are stored in MongoDB and synced via authenticated API calls.', style='List Bullet')
doc.add_paragraph('• Bookmark actions refresh recommendations automatically.', style='List Bullet')
doc.add_paragraph('• Profile updates are saved via PUT /api/auth/profile and user objects are persisted in localStorage.', style='List Bullet')

doc.add_heading('5. Deployment', level=1)
doc.add_paragraph('The application can be deployed with separate frontend and backend hosting. The frontend is suitable for Vercel or Netlify, while the backend runs on a Python host such as Render or any other containerized service.')

doc.add_heading('5.1 Frontend Deployment', level=2)
doc.add_paragraph('The ui/ folder is the frontend source. Vercel deploys it as a static Angular app. Key configuration:')
doc.add_paragraph('• ui/vercel.json includes an SPA rewrite to index.html.', style='List Bullet')
doc.add_paragraph('• ui/package.json contains the production build script and vercel-build command.', style='List Bullet')
doc.add_paragraph('• ui/src/environments/environment.prod.ts points to the production backend API URL.', style='List Bullet')

doc.add_heading('5.2 Backend Deployment', level=2)
doc.add_paragraph('The backend uses Uvicorn and requires Python dependencies installed from requirements.txt. Ensure environment variables are set for MongoDB, JWT secret, and Google client ID.', style='List Bullet')
doc.add_paragraph('• Use either repo/requirements.txt or api/requirements.txt depending on your host configuration.', style='List Bullet')

doc.add_heading('6. Local Setup and Run', level=1)
doc.add_paragraph('To run locally:')
doc.add_paragraph('1. Backend setup', style='List Number')
doc.add_paragraph('   • Create a Python virtual environment and activate it.', style='List Bullet')
doc.add_paragraph('   • Install backend dependencies with pip.', style='List Bullet')
doc.add_paragraph('   • Create a .env file containing MONGODB_URL, JWT_SECRET, and GOOGLE_CLIENT_ID.', style='List Bullet')
doc.add_paragraph('   • Start the backend with uvicorn main:app --reload.', style='List Bullet')
doc.add_paragraph('2. Frontend setup', style='List Number')
doc.add_paragraph('   • Navigate to ui/ and run npm install.', style='List Bullet')
doc.add_paragraph('   • Start the frontend with npm start or npx ng serve.', style='List Bullet')
doc.add_paragraph('3. Test the application', style='List Number')
doc.add_paragraph('   • Open http://localhost:4200 and sign in with Google.', style='List Bullet')

doc.add_heading('7. Key Files Summary', level=1)
doc.add_paragraph('Backend:')
doc.add_paragraph('• api/main.py: FastAPI endpoints, CORS, MongoDB connection, auth flow.', style='List Bullet')
doc.add_paragraph('• api/requirements.txt: Backend dependencies.', style='List Bullet')
doc.add_paragraph('• api/scripts/insert_opportunities.py: Opportunity seeding script.', style='List Bullet')

doc.add_paragraph('Frontend:')
doc.add_paragraph('• ui/src/app/services/auth.service.ts: Auth and bookmark API integration.', style='List Bullet')
doc.add_paragraph('• ui/src/app/components/login/login.component.ts: Google login component.', style='List Bullet')
doc.add_paragraph('• ui/src/app/components/recommendations/recommendations.component.ts: Recommendation generation.', style='List Bullet')
doc.add_paragraph('• ui/src/environments/environment.prod.ts: Production backend URL.', style='List Bullet')

doc.add_heading('8. Notes and Next Steps', level=1)
doc.add_paragraph('• The recommendation logic can be improved with richer user preferences and more advanced scoring.', style='List Bullet')
doc.add_paragraph('• Production deployment should secure JWT secrets and MongoDB credentials, and restrict CORS origins.', style='List Bullet')
doc.add_paragraph('• Add validation and error handling on both frontend and backend for stronger robustness.', style='List Bullet')

filename = 'ScholarSite_Technical_Documentation.docx'
doc.save(filename)
print(filename)
